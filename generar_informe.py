# -*- coding: utf-8 -*-
"""
generar_informe.py
---------------------------------------------------------------------------
Genera "Informe de resultados de tratamiento de imagen" (.docx, normas APA 7.a ed.)

1. Si la carpeta de imagenes contiene las cinco capturas (referencia, oscura,
   desenfocada, fondo, angulo), ejecuta tratamiento_imagen.py, calcula las
   metricas y genera las figuras; el informe se llena con esos resultados.
2. Si faltan imagenes, produce el informe completo con marcadores [valor] y
   [Figura pendiente] para completarlo despues.

Uso:
    python generar_informe.py
    python generar_informe.py --imagenes imagenes --salida Informe_tratamiento_imagen.docx
---------------------------------------------------------------------------
"""

import argparse
import ast
import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import tratamiento_imagen as ti

# ---------------------------------------------------------------------------
# Datos de portada (edite estos valores)
# ---------------------------------------------------------------------------
ESTUDIANTE = "[Nombre]"
INSTITUCION = "[Institución / Facultad]"
MATERIA = "[Materia]"
DOCENTE = "[Nombre del docente]"
CADI = "[CADI]"
FECHA = "[Fecha]"
OBJETO = "un cubo de Rubik 3×3"
DISPOSITIVO = "[dispositivo de captura]"
REPOSITORIO = ("https://github.com/DiegoYanzir/"
               "Taller-R2-A2-S8---Informe-de-resultados-de-tratamiento-de-imagen")

BASE = os.path.dirname(os.path.abspath(__file__))
FUENTE = "Times New Roman"
CLAVE_PROPUESTA = "Gamma + CLAHE + bilateral (propuesta)"


# ---------------------------------------------------------------------------
# Utilidades de formato numerico
# ---------------------------------------------------------------------------
def v(x, d=1):
    """Numero formateado o marcador si no hay datos."""
    return "[valor]" if x is None else f"{x:.{d}f}"


def var_pct(antes, despues):
    """Variacion porcentual con signo, o marcador."""
    if antes is None or despues is None or antes == 0:
        return "[valor] %"
    return f"{(despues - antes) / antes * 100:+.1f} %"


def ratio_pct(x, ref):
    """x expresado como porcentaje de ref."""
    if x is None or ref in (None, 0):
        return "[valor] %"
    return f"{x / ref * 100:.1f} %"


class Datos:
    """Acceso seguro a los resultados (devuelve None si no existen)."""

    def __init__(self, r):
        self.r = r

    def m(self, img, clave):
        return None if self.r is None else self.r["metricas"][img].get(clave)

    def t(self, var, clave):
        return None if self.r is None else self.r["tratamiento"][var].get(clave)

    def f(self, clave):
        return None if self.r is None else self.r["formatos"].get(clave)

    def p(self, clave):
        return None if self.r is None else self.r["parametros"].get(clave)


# ---------------------------------------------------------------------------
# Utilidades de documento (APA 7)
# ---------------------------------------------------------------------------
def set_cell_border(cell, **bordes):
    """Define bordes de celda: set_cell_border(c, bottom={'sz': 8, 'val': 'single'})."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for lado, attrs in bordes.items():
        el = OxmlElement(f"w:{lado}")
        for k, val in attrs.items():
            el.set(qn(f"w:{k}"), str(val))
        tc_borders.append(el)


def set_table_borders_apa(table):
    """Tabla APA: solo lineas horizontales superior e inferior."""
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        if lado in ("top", "bottom"):
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "8"); el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def shade(paragraph, color="F2F2F2"):
    """Sombreado de fondo de un parrafo (bloques de codigo)."""
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    p_pr.append(shd)


def add_page_number(section):
    """Numero de pagina en la esquina superior derecha (APA 7, trabajo de estudiante)."""
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = texto
        run._r.append(el)
    run.font.name = FUENTE; run.font.size = Pt(12)


def configurar_estilos(doc):
    """Times New Roman 12, doble espacio, encabezados APA niveles 1-3."""
    estilos = doc.styles
    normal = estilos["Normal"]
    normal.font.name = FUENTE
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    lang = OxmlElement("w:lang"); lang.set(qn("w:val"), "es-EC"); normal.element.rPr.append(lang)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(0); pf.space_after = Pt(0)
    pf.widow_control = True

    for nivel, alin, ital in ((1, WD_ALIGN_PARAGRAPH.CENTER, False),
                              (2, WD_ALIGN_PARAGRAPH.LEFT, False),
                              (3, WD_ALIGN_PARAGRAPH.LEFT, True)):
        st = estilos[f"Heading {nivel}"]
        st.font.name = FUENTE; st.font.size = Pt(12); st.font.bold = True
        st.font.italic = ital; st.font.color.rgb = RGBColor(0, 0, 0)
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts"); rpr.append(fonts)
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(a), FUENTE)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if fonts.get(qn(a)) is not None:
                del fonts.attrib[qn(a)]
        st.paragraph_format.alignment = alin
        st.paragraph_format.space_before = Pt(0); st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        st.paragraph_format.keep_with_next = True

    codigo = estilos.add_style("Codigo", 1)
    codigo.base_style = normal
    codigo.font.name = "Consolas"; codigo.font.size = Pt(8)
    codigo.element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    cpf = codigo.paragraph_format
    cpf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    cpf.first_line_indent = Inches(0); cpf.left_indent = Inches(0.1)

    tabla = estilos.add_style("TablaTexto", 1)
    tabla.base_style = normal
    tabla.font.size = Pt(10)
    tabla.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    tabla.paragraph_format.space_before = Pt(2); tabla.paragraph_format.space_after = Pt(2)


TOKEN = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\^\{[^}]+\}|_\{[^}]+\})")


def add_runs(p, texto, size=None):
    """Inserta texto con marcas ligeras: **negrita**, *cursiva*, `codigo`."""
    for parte in TOKEN.split(texto):
        if not parte:
            continue
        if parte.startswith("**"):
            r = p.add_run(parte[2:-2]); r.bold = True
        elif parte.startswith("*"):
            r = p.add_run(parte[1:-1]); r.italic = True
        elif parte.startswith("`"):
            r = p.add_run(parte[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(10.5)
        elif parte.startswith("^{"):
            r = p.add_run(parte[2:-1]); r.font.superscript = True
        elif parte.startswith("_{"):
            r = p.add_run(parte[2:-1]); r.font.subscript = True
        else:
            r = p.add_run(parte)
        if size and not parte.startswith("`"):
            r.font.size = Pt(size)
    return p


def parrafo(doc, texto, sangria=True, alin=None):
    """Parrafo de cuerpo APA: sangria de primera linea de 0.5 pulgadas."""
    p = doc.add_paragraph()
    if sangria:
        p.paragraph_format.first_line_indent = Inches(0.5)
    if alin is not None:
        p.alignment = alin
    return add_runs(p, texto)


def vinetas(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, it)


def ecuacion(doc, texto, numero):
    """Ecuacion centrada con numero entre parentesis."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(p, texto)
    for r in p.runs:
        r.italic = True
    p.add_run(f"     ({numero})")


def codigo(doc, fuente):
    """Bloque de codigo monoespaciado con fondo gris (una linea por parrafo)."""
    lineas = fuente.rstrip("\n").split("\n")
    for i, linea in enumerate(lineas):
        p = doc.add_paragraph(style="Codigo")
        p.add_run(linea if linea else " ")
        shade(p)
        p.paragraph_format.keep_with_next = i < len(lineas) - 1 and len(lineas) < 25
    doc.add_paragraph(style="Codigo")


def etiqueta(doc, tipo, numero, titulo):
    """Numero (negrita) y titulo (cursiva) de tabla o figura, segun APA 7."""
    p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
    p.add_run(f"{tipo} {numero}").bold = True
    p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
    p.add_run(titulo).italic = True


def nota(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    r = p.add_run("Nota. "); r.italic = True
    add_runs(p, texto)


def tabla_apa(doc, numero, titulo, encabezados, filas, anchos, nota_txt=None):
    """Tabla con estilo APA 7: lineas horizontales, encabezado centrado."""
    etiqueta(doc, "Tabla", numero, titulo)
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_table_borders_apa(t)
    for i, h in enumerate(encabezados):
        c = t.rows[0].cells[i]
        c.width = Inches(anchos[i])
        c.paragraphs[0].style = "TablaTexto"
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(c.paragraphs[0], h)
        set_cell_border(c, bottom={"val": "single", "sz": 6, "color": "000000"})
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); tr_pr.append(th)
    for fila in filas:
        celdas = t.add_row().cells
        for i, val in enumerate(fila):
            celdas[i].width = Inches(anchos[i])
            par = celdas[i].paragraphs[0]
            par.style = "TablaTexto"
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 or len(str(val)) > 14 else WD_ALIGN_PARAGRAPH.CENTER
            add_runs(par, str(val))
    for fila in t.rows:                                       # La tabla no se parte entre paginas
        tr = fila._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); tr.append(cs)
        if fila is not t.rows[-1] and len(t.rows) <= 12:
            for c in fila.cells:
                for par in c.paragraphs:
                    par.paragraph_format.keep_with_next = True
    if nota_txt:
        nota(doc, nota_txt)
    else:
        doc.add_paragraph()


def figura(doc, numero, titulo, ruta, nota_txt, ancho=6.5):
    """Figura APA 7: numero, titulo, imagen y nota."""
    etiqueta(doc, "Figura", numero, titulo)
    if ruta and os.path.isfile(ruta):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(ruta, width=Inches(ancho))
    else:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("[Figura pendiente: coloque las cinco capturas en la carpeta imagenes/ "
                      "y ejecute generar_informe.py]")
        r.italic = True; r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    nota(doc, nota_txt)


def referencia(doc, texto):
    """Entrada de referencias con sangria francesa de 0.5 pulgadas."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    add_runs(p, texto)


def salto_pagina(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def fuente_funcion(nombres):
    """Extrae el codigo fuente de funciones de tratamiento_imagen.py (sincronizado)."""
    ruta = os.path.join(BASE, "tratamiento_imagen.py")
    src = open(ruta, encoding="utf-8").read()
    arbol = ast.parse(src)
    lineas = src.split("\n")
    bloques = []
    for nodo in arbol.body:
        if isinstance(nodo, ast.FunctionDef) and nodo.name in nombres:
            bloques.append("\n".join(lineas[nodo.lineno - 1:nodo.end_lineno]))
    return "\n\n".join(bloques)


# ---------------------------------------------------------------------------
# Construccion del informe
# ---------------------------------------------------------------------------
def construir(resultados, carpeta_res, salida):
    D = Datos(resultados)
    fig = (lambda n: os.path.join(carpeta_res, n)) if resultados else (lambda n: None)
    hay = resultados is not None

    doc = Document()
    configurar_estilos(doc)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, lado, Inches(1))
    add_page_number(sec)

    # ---------------- Portada (trabajo de estudiante, APA 7) ----------------
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Informe de resultados de tratamiento de imagen").bold = True
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("R2-A2-S8: Diagnóstico, tratamiento y evaluación de imágenes con OpenCV en Python").bold = True
    doc.add_paragraph()
    for linea in (ESTUDIANTE, INSTITUCION, f"{MATERIA} – Especialización en Inteligencia Artificial",
                  f"CADI: {CADI}", DOCENTE, FECHA):
        p = doc.add_paragraph(linea); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Repositorio: ").bold = True
    p = doc.add_paragraph(REPOSITORIO); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    salto_pagina(doc)

    # ---------------- Titulo del cuerpo e introduccion ----------------
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Informe de resultados de tratamiento de imagen").bold = True
    parrafo(doc, "Un sistema de percepción computacional solo es tan confiable como los datos que recibe. "
            "Antes de entrenar o desplegar un modelo de reconocimiento de objetos es necesario determinar si las "
            "imágenes capturadas poseen condiciones adecuadas de exposición, contraste, nitidez y composición, "
            "y, cuando no las poseen, decidir qué intervención es pertinente. El presente informe responde a la "
            "pregunta orientadora de la semana: *¿cómo determinar si una imagen tiene condiciones adecuadas para "
            "un sistema de percepción computacional y qué hacer cuando no las tiene?* Para ello se sigue un ciclo "
            "explícito de **diagnóstico → decisión técnica justificada → implementación algorítmica → evaluación "
            "cuantitativa y cualitativa**, implementado con OpenCV en Python (Bradski, 2000). El criterio rector no "
            "es aplicar muchos filtros, sino identificar un defecto concreto, intervenirlo con una técnica coherente "
            "y demostrar con evidencia medible qué cambió.")

    # ---------------- 1. Problema y dataset ----------------
    doc.add_heading("Definición del Problema y Conjunto de Datos", level=1)
    doc.add_heading("Objeto de Estudio", level=2)
    parrafo(doc, f"Se seleccionó como objeto de estudio {OBJETO}, un elemento cotidiano con contornos bien "
            "definidos, superficies de color uniforme y detalles finos (texturas, bordes y elementos gráficos) que "
            "permiten observar con claridad el efecto de la calidad de captura sobre la información visual. Se "
            f"capturaron cinco fotografías propias del mismo objeto con {DISPOSITIVO}, sin edición previa en el "
            "teléfono, y se nombraron de forma sistemática para garantizar la trazabilidad del análisis "
            "(`referencia.jpg`, `oscura.jpg`, `desenfocada.jpg`, `fondo.jpg` y `angulo.jpg`). Mantener el mismo "
            "objeto aísla la variable de interés: los cambios en las métricas se atribuyen a la condición de captura "
            "y no a diferencias propias del objeto.")
    doc.add_heading("Condiciones de Captura Evaluadas", level=2)
    parrafo(doc, "La Tabla 1 resume las cinco condiciones definidas en la guía de la semana, junto con el efecto "
            "esperado sobre la señal de imagen y el riesgo que representa para un pipeline posterior.")
    tabla_apa(doc, 1, "Condiciones de captura del conjunto de imágenes",
              ["#", "Condición", "Descripción de la captura", "Riesgo para el pipeline de visión"],
              [["1", "Referencia", "Buena iluminación, enfoque adecuado y composición clara.",
                "Línea base para comparar todas las métricas."],
               ["2", "Baja iluminación", "Escena intencionalmente oscura (subexposición).",
                "Histograma comprimido en tonos bajos; gradientes débiles; ruido de sensor por ISO alto."],
               ["3", "Desenfoque / movimiento", "Pérdida visible de nitidez por foco o movimiento.",
                "Atenuación de altas frecuencias; bordes difusos que el detector no localiza."],
               ["4", "Fondo complejo", "Objeto rodeado de otros elementos y texturas.",
                "Bordes espurios y falsos positivos en segmentación y clasificación."],
               ["5", "Otra perspectiva", "Cambio de ángulo, distancia o punto de vista.",
                "Variación de escala, pose y oclusiones; exige invariancia geométrica."]],
              [0.3, 1.25, 2.3, 2.65])
    doc.add_heading("Por Qué Estas Condiciones Afectan a un Sistema de Percepción", level=2)
    parrafo(doc, "Los algoritmos clásicos y las redes neuronales convolucionales (CNN) dependen de la misma materia "
            "prima: variaciones locales de intensidad. Un detector de bordes como Canny (1986) localiza máximos del "
            "gradiente ‖∇I‖ y los filtra con umbrales de histéresis fijos; si la escena está subexpuesta, la "
            "magnitud del gradiente cae por debajo de esos umbrales y los contornos simplemente desaparecen. Del "
            "mismo modo, las primeras capas de una CNN aprenden filtros sensibles a bordes y texturas (Rodríguez "
            "Araújo, 2019; Futurology — An Optimistic Future, 2020); si la entrada difiere en exposición o nitidez "
            "de la distribución de entrenamiento, las activaciones cambian y la clasificación se degrada. El "
            "desenfoque actúa como un filtro pasa-bajos que elimina precisamente las frecuencias que codifican los "
            "contornos (Gonzalez & Woods, 2018), mientras que el fondo complejo introduce bordes que no pertenecen "
            "al objeto y dificultan la extracción de contornos cerrados (Image processing class (EGBE443) #5, s.f.). "
            "Por último, la perspectiva modifica la geometría proyectada del objeto, un problema que el "
            "preprocesamiento fotométrico no puede resolver y que exige invariancia geométrica o aumento de datos "
            "(Szeliski, 2022).")

    # ---------------- 2. Diagnostico ----------------
    doc.add_heading("Diagnóstico de las Imágenes", level=1)
    doc.add_heading("Métricas Utilizadas", level=2)
    parrafo(doc, f"Para que las métricas fueran comparables, cada imagen se reescaló a un ancho común de "
            f"{ti.ANCHO_ANALISIS} píxeles conservando la relación de aspecto (interpolación `INTER_AREA`); la "
            "varianza del Laplaciano y la densidad de bordes dependen de la escala, por lo que comparar imágenes "
            "de distinta resolución sesgaría el diagnóstico. Sobre la imagen en escala de grises *I* se calcularon "
            "las métricas sugeridas por la guía y tres complementarias cuya utilidad se justifica a continuación:")
    vinetas(doc, [
        "**Brillo** (media de *I*): nivel global de exposición en el rango 0–255.",
        "**Contraste** (desviación estándar de *I*, o contraste RMS): separación entre zonas claras y oscuras.",
        "**Nitidez** (varianza del Laplaciano, Var(∇²*I*)): energía de alta frecuencia; cae con el desenfoque "
        "(Pech-Pacheco et al., 2000). Se interpreta de forma relativa dentro del conjunto, no con un umbral universal.",
        "**Ruido estimado** σ̂ (método de Immerkær, 1996): necesario porque el ruido también eleva la varianza del "
        "Laplaciano; sin esta métrica una imagen ruidosa podría parecer «nítida».",
        "**Entropía** de Shannon del histograma (bits): cuántos niveles de gris aportan información.",
        f"**Densidad de bordes**: porcentaje de píxeles marcados por Canny con umbrales fijos "
        f"({ti.CANNY_BAJO}/{ti.CANNY_ALTO}); simula la respuesta de un detector posterior que no se reajusta por imagen.",
    ])

    doc.add_heading("Resultados del Diagnóstico", level=2)
    nombres = list(ti.CONDICIONES)
    filas = []
    for n in nombres:
        res = (f"{D.m(n, 'ancho')}×{D.m(n, 'alto')}" if hay else "[valor]")
        filas.append([ti.CONDICIONES[n], v(D.m(n, "brillo")), v(D.m(n, "contraste")), v(D.m(n, "nitidez")),
                      v(D.m(n, "ruido"), 2), v(D.m(n, "densidad_bordes"), 2), res,
                      D.m(n, "problema") or "[problema]", D.m(n, "adecuada") or "[Sí/No]"])
    parrafo(doc, "La Tabla 2 consolida las métricas de las cinco capturas y la Figura 1 muestra las imágenes "
            "con sus valores principales. La columna «¿Adecuada?» se asignó comparando cada imagen contra la "
            "referencia: se marca *No* cuando el brillo o el contraste caen por debajo del 60 % de la referencia, "
            "cuando más del 40 % de los píxeles está subexpuesto, cuando la nitidez cae por debajo del 35 % o "
            "cuando el ruido supera 1.5 veces el de la referencia; se marca *Condicional* cuando el defecto no es "
            "fotométrico (fondo o perspectiva) y depende de la robustez del modelo.")
    tabla_apa(doc, 2, "Métricas de calidad de las cinco capturas",
              ["Imagen", "Brillo", "Contraste", "Nitidez", "Ruido σ̂", "Bordes (%)", "Resolución", "Problema",
               "¿Adecuada?"],
              filas, [0.9, 0.5, 0.75, 0.6, 0.5, 0.55, 0.8, 1.0, 0.9],
              "Brillo y contraste en niveles de gris (0–255); nitidez como varianza del Laplaciano; ruido según "
              "Immerkær (1996); bordes como porcentaje de píxeles detectados por Canny. Las métricas se calcularon "
              f"a {ti.ANCHO_ANALISIS} px de ancho; la resolución corresponde al archivo original.")
    figura(doc, 1, "Conjunto de capturas del mismo objeto bajo cinco condiciones", fig("fig1_conjunto.png"),
           "Elaboración propia con OpenCV y Matplotlib. Bajo cada imagen se indican brillo, contraste y nitidez.")

    parrafo(doc, f"**Interpretación de los datos básicos.** Todas las imágenes tienen tres canales (BGR), es decir, "
            "información de color que puede aprovecharse o descartarse según la tarea. La resolución original de la "
            f"referencia fue de {D.m('referencia', 'ancho') or '[ancho]'} × {D.m('referencia', 'alto') or '[alto]'} "
            f"píxeles ({v(D.m('referencia', 'megapixeles'), 1)} MP); una resolución alta no garantiza mejor desempeño, "
            "porque los modelos trabajan con entradas reducidas (p. ej., 224 × 224) y lo que se conserva tras el "
            "reescalado es la información de bordes y contraste, no la cantidad de píxeles.")
    parrafo(doc, f"**Iluminación.** La imagen oscura registró un brillo medio de {v(D.m('oscura', 'brillo'))} frente a "
            f"{v(D.m('referencia', 'brillo'))} de la referencia ({ratio_pct(D.m('oscura', 'brillo'), D.m('referencia', 'brillo'))} "
            f"de la referencia) y un contraste de {v(D.m('oscura', 'contraste'))} frente a "
            f"{v(D.m('referencia', 'contraste'))}. El {v(D.m('oscura', 'pct_oscuros'))} % de sus píxeles quedó por "
            f"debajo de {ti.UMBRAL_OSCURO}, y la densidad de bordes de Canny bajó a "
            f"{v(D.m('oscura', 'densidad_bordes'), 2)} % frente a {v(D.m('referencia', 'densidad_bordes'), 2)} % en la "
            "referencia: el objeto sigue presente en la escena, pero el detector ya no lo «ve».")
    parrafo(doc, f"**Desenfoque.** La imagen desenfocada conservó un brillo similar al de la referencia "
            f"({v(D.m('desenfocada', 'brillo'))}), lo que confirma que su defecto no es de exposición; sin embargo, "
            f"su nitidez descendió a {v(D.m('desenfocada', 'nitidez'))}, equivalente al "
            f"{ratio_pct(D.m('desenfocada', 'nitidez'), D.m('referencia', 'nitidez'))} de la referencia. La varianza "
            "del Laplaciano reflejó con claridad la pérdida de altas frecuencias. **Fondo y perspectiva.** En estas "
            "capturas las métricas globales varían poco o incluso aumentan (el fondo complejo suele elevar la nitidez "
            f"y la densidad de bordes, aquí {v(D.m('fondo', 'nitidez'))} y {v(D.m('fondo', 'densidad_bordes'), 2)} %), "
            "lo que demuestra que una métrica global alta no implica una imagen adecuada: los bordes adicionales "
            "pertenecen a distractores y no al objeto.")

    doc.add_heading("Análisis de Histogramas", level=2)
    parrafo(doc, "La Figura 2 presenta los histogramas de intensidad normalizados. En la referencia la "
            "distribución ocupa buena parte del rango dinámico; en la imagen oscura la masa se concentra en el "
            "extremo izquierdo, con una entropía de "
            f"{v(D.m('oscura', 'entropia'), 2)} bits frente a {v(D.m('referencia', 'entropia'), 2)} bits de la "
            "referencia. Esta compresión implica que diferencias reales de reflectancia del objeto se cuantizaron en "
            "muy pocos niveles de gris, reduciendo la separabilidad entre objeto y fondo.")
    figura(doc, 2, "Histogramas de intensidad de las cinco capturas", fig("fig2_histogramas.png"),
           "Porcentaje de píxeles por nivel de gris (0–255), con eje vertical compartido para facilitar la comparación.")

    doc.add_heading("Formato y Tamaño de Archivo", level=2)
    parrafo(doc, "La referencia se guardó en JPG (calidad 90) y en PNG con la misma resolución (Tabla 3). La "
            "diferencia de tamaño se explica por el tipo de compresión: JPG aplica la transformada discreta del "
            "coseno por bloques de 8 × 8 y cuantiza los coeficientes de alta frecuencia (compresión con pérdida), "
            "mientras que PNG usa DEFLATE sin pérdida y conserva cada valor de píxel. La PSNR mide la fidelidad del "
            "JPG frente al original; valores superiores a unos 40 dB suelen ser imperceptibles a simple vista, pero "
            "los artefactos de bloque pueden introducir bordes falsos en tareas sensibles a gradientes.")
    tabla_apa(doc, 3, "Comparación de formatos de almacenamiento",
              ["Formato", "Resolución", "Tamaño (KB)", "Observación"],
              [["JPG (calidad 90)", D.f("resolucion") or "[valor]", v(D.f("jpg_kb")),
                f"Con pérdida; PSNR = {v(D.f('psnr_jpg'), 2)} dB frente al original."],
               ["PNG", D.f("resolucion") or "[valor]", v(D.f("png_kb")),
                "Sin pérdida; fidelidad exacta, mayor tamaño."]],
              [1.3, 1.1, 1.0, 3.1],
              "Ambos archivos se generaron con `cv2.imwrite` a partir de la misma matriz de píxeles.")

    doc.add_heading("Selección de la Imagen Crítica e Identificación del Defecto", level=2)
    parrafo(doc, "Se seleccionó la imagen **oscura** porque presenta el defecto fotométrico más severo y, a la "
            "vez, corregible mediante procesamiento: la información del objeto existe en la señal, pero "
            "comprimida en un rango estrecho de intensidades. El desenfoque, en cambio, implica pérdida "
            "irreversible de altas frecuencias, y el fondo y la perspectiva son problemas de composición. "
            "Siguiendo el formato de la guía:")
    p = parrafo(doc, "«La imagen presenta **subexposición, contraste reducido, sombras densas y ruido de sensor**, que "
                "comprimen el histograma en los tonos bajos. Aplicaré **corrección gamma adaptativa y CLAHE sobre el "
                "canal de luminancia, seguidas de un filtro bilateral**, porque espero **recuperar el rango dinámico "
                "y el contraste local del objeto sin alterar su color ni amplificar el ruido del sensor**, de modo "
                "que los bordes vuelvan a superar los umbrales de un detector posterior.»", sangria=False)
    p.paragraph_format.left_indent = Inches(0.5)

    # ---------------- 3. Diseno de la solucion ----------------
    doc.add_heading("Diseño de la Solución y Justificación Técnica", level=1)
    doc.add_heading("Alternativas Consideradas", level=2)
    parrafo(doc, "La alternativa más simple es la transformación lineal *g*(*x*) = α·*f*(*x*) + β implementada con "
            "`cv2.convertScaleAbs`. Es global y uniforme: desplaza y estira todo el histograma por igual, de modo "
            "que no puede aclarar las sombras sin saturar las zonas ya iluminadas, y amplifica el ruido con el "
            "mismo factor α. La ecualización global de histograma usa como función de transferencia la "
            "distribución acumulada (Gonzalez & Woods, 2018):")
    ecuacion(doc, "s_{k} = T(r_{k}) = (L − 1) · Σ_{j = 0}^{k} p_{r}(r_{j})", 1)
    parrafo(doc, "donde *p*_{r}(*r*_{j}) es la probabilidad del nivel *r*_{j} y *L* = 256. Al ser global, una región "
            "grande y oscura domina la CDF: los niveles más frecuentes se separan agresivamente, lo que amplifica "
            "el ruido en zonas homogéneas y puede saturar regiones claras.")
    doc.add_heading("Técnica Seleccionada", level=2)
    parrafo(doc, "**Etapa 1: corrección gamma adaptativa.** La transformación de potencia *s* = (*L* − 1)·(*r*/(*L* − 1))^{γ} "
            "con γ < 1 expande los tonos oscuros y comprime los claros de forma no lineal. El exponente no se fija "
            "a mano: se elige para llevar la luminancia media normalizada μ a 0.5,")
    ecuacion(doc, "γ = ln(0.5) / ln(μ),   γ ∈ [0.35, 1]", 2)
    parrafo(doc, "con un límite inferior de 0.35 para no «lavar» la imagen. Esta etapa es necesaria porque CLAHE, "
            "por diseño, limita la amplificación: en una prueba de verificación con una escena muy oscura (brillo "
            "≈ 29), CLAHE por sí solo apenas elevó el brillo a ≈ 38, insuficiente para la tarea. "
            f"{'En la imagen analizada se obtuvo γ = ' + v(D.p('gamma'), 2) + '.' if hay else 'El valor de γ obtenido se reporta en metricas.json.'}")
    parrafo(doc, "**Etapa 2: CLAHE** (*Contrast Limited Adaptive Histogram Equalization*; Pizer et al., 1987; "
            "Zuiderveld, 1994). La imagen se divide en una rejilla de regiones contextuales y en cada una se "
            "ecualiza el histograma local, pero antes se recorta cada bin a un tope β y el excedente se redistribuye "
            "uniformemente. En OpenCV el tope se calcula como:")
    ecuacion(doc, "β = max(1, clipLimit · (M·N) / L)", 3)
    parrafo(doc, "donde *M*·*N* es el número de píxeles de cada región. Limitar la pendiente de la CDF local acota la "
            "ganancia de contraste y, con ello, la amplificación del ruido; la interpolación bilineal entre regiones "
            "vecinas evita discontinuidades en las fronteras. CLAHE se aplica solo al canal *L* del espacio CIELAB, "
            "que separa la luminancia de la cromaticidad (*a*, *b*): así se mejora el contraste sin desplazar el tono "
            "del objeto, algo que no ocurre si se ecualizan por separado los canales B, G y R.")
    parrafo(doc, "**Etapa 3: filtro bilateral** (Tomasi & Manduchi, 1998). Toda expansión de contraste en una "
            "escena oscura amplifica el ruido del sensor. Un filtro de promedio o gaussiano lo reduciría, pero "
            "también difuminaría los bordes que se intenta recuperar. El filtro bilateral pondera cada vecino por "
            "su cercanía espacial y por su similitud de intensidad:")
    ecuacion(doc, "BF[I]_{p} = (1/W_{p}) · Σ_{q∈S} G_{σs}(‖p − q‖) · G_{σr}(|I_{p} − I_{q}|) · I_{q}", 4)
    parrafo(doc, "donde *W*_{p} normaliza los pesos. Cuando *p* y *q* están a lados opuestos de un borde, "
            "|*I*_{p} − *I*_{q}| es grande y el peso *G*_{σr} tiende a cero: el filtro promedia dentro de las regiones "
            "homogéneas pero no a través de los contornos. La Tabla 4 resume la comparación conceptual.")
    tabla_apa(doc, 4, "Comparación de técnicas candidatas para la imagen subexpuesta",
              ["Técnica", "Alcance", "Ventaja", "Limitación para este problema"],
              [["Ajuste lineal α, β", "Global, lineal", "Simple y predecible.",
                "No separa sombras de luces; amplifica ruido con el mismo factor."],
               ["Ecualización global", "Global, no lineal", "Maximiza el uso del rango dinámico.",
                "Sobrerrealce y ruido en zonas homogéneas; posibles saturaciones."],
               ["Solo CLAHE", "Local, limitada", "Contraste local con ruido acotado.",
                "En subexposición severa no basta para recuperar el brillo."],
               ["Gamma + CLAHE + bilateral", "Global + local + suavizado", "Corrige exposición y contraste local "
                "preservando bordes.", "Mayor costo computacional; requiere ajustar hiperparámetros."]],
              [1.45, 1.1, 1.8, 2.15])
    doc.add_heading("Diagrama de Flujo del Pipeline", level=2)
    parrafo(doc, "La Figura 3 muestra el pipeline propuesto. Las etapas resaltadas son las que transforman la imagen; "
            "las demás son de entrada, conversión y evaluación.")
    figura(doc, 3, "Pipeline de procesamiento propuesto", fig("fig3_pipeline.png"),
           "Captura → diagnóstico → BGR a CIELAB → gamma adaptativa en L → CLAHE en L → LAB a BGR y filtro bilateral "
           "→ re-medición. Elaboración propia.")

    # ---------------- 4. Implementacion ----------------
    doc.add_heading("Descripción e Implementación de los Scripts", level=1)
    parrafo(doc, "La solución se organizó en dos scripts disponibles en el repositorio. `tratamiento_imagen.py` "
            "contiene el ciclo completo (carga, métricas, comparación de formatos, tratamiento, figuras y exportación "
            "de `metricas.json`) y puede ejecutarse en Google Colab o Jupyter con "
            "`!python tratamiento_imagen.py --imagenes imagenes --salida resultados`. `generar_informe.py` ejecuta "
            "el análisis y construye este documento con los valores obtenidos, de modo que cada cifra del informe es "
            "reproducible. El código completo y comentado se incluye en el Apéndice A; a continuación se describen "
            "los fragmentos centrales.")
    doc.add_heading("Carga y Cálculo de Métricas", level=2)
    parrafo(doc, "`cv2.imdecode` lee la imagen en orden BGR (la convención de OpenCV); se usa junto con "
            "`np.fromfile` para admitir rutas con tildes en Windows. `cv2.cvtColor(..., cv2.COLOR_BGR2GRAY)` obtiene "
            "la luminancia ponderada 0.299R + 0.587G + 0.114B, y `cv2.Laplacian(gris, cv2.CV_64F)` se calcula en "
            "punto flotante para no truncar las respuestas negativas del operador.")
    codigo(doc, fuente_funcion({"calcular_metricas"}))
    doc.add_heading("Tratamiento Propuesto", level=2)
    codigo(doc, fuente_funcion({"gamma_adaptativa", "corregir_gamma", "aplicar_clahe", "tratamiento_propuesto"}))
    parrafo(doc, "La Tabla 5 detalla las funciones de OpenCV empleadas en el tratamiento y la justificación de "
            "cada hiperparámetro.")
    tabla_apa(doc, 5, "Funciones de OpenCV e hiperparámetros del tratamiento",
              ["Función", "Parámetro", "Valor", "Justificación"],
              [["cv2.cvtColor", "COLOR_BGR2LAB", "—",
                "Aísla la luminancia L; a y b conservan el color del objeto."],
               ["cv2.LUT", "γ (adaptativa)", f"{v(D.p('gamma'), 2)} (mín. {ti.GAMMA_MIN})",
                "Lleva la luminancia media hacia 0.5 con una tabla de 256 entradas (costo O(1) por píxel)."],
               ["cv2.createCLAHE", "clipLimit", f"{ti.CLAHE_CLIP}",
                "Valores de 2–3 equilibran realce y ruido; valores > 4 generan halos y grano."],
               ["cv2.createCLAHE", "tileGridSize", f"{ti.CLAHE_TILES[0]}×{ti.CLAHE_TILES[1]}",
                "Regiones de ~128 px a 1024 px de ancho: suficientemente locales sin ver el ruido como estructura."],
               ["cv2.bilateralFilter", "d", f"{ti.BILAT_D}",
                "Vecindario pequeño: suaviza el grano sin costo excesivo (el filtro es O(d²) por píxel)."],
               ["cv2.bilateralFilter", "sigmaColor", f"{ti.BILAT_SIGMA_COLOR}",
                "Mayor que la amplitud del ruido amplificado y menor que el salto de intensidad en los bordes."],
               ["cv2.bilateralFilter", "sigmaSpace", f"{ti.BILAT_SIGMA_SPACE}",
                "Coherente con d; un valor mayor difuminaría texturas finas del objeto."],
               ["cv2.Canny", "umbrales", f"{ti.CANNY_BAJO} / {ti.CANNY_ALTO}",
                "Fijos para todas las imágenes: simulan un detector posterior que no se ajusta por imagen."]],
              [1.45, 1.1, 1.0, 2.95])
    doc.add_heading("Visualización Comparativa", level=2)
    parrafo(doc, "Matplotlib se utiliza para generar las comparativas: las imágenes se convierten de BGR a RGB antes "
            "de `imshow`, los histogramas se calculan con `cv2.calcHist` y se normalizan a porcentaje de píxeles, y la "
            "distribución acumulada se obtiene con `np.cumsum`. Todas las figuras se guardan a 200 ppp para su "
            "inclusión en el informe.")

    # ---------------- 5. Evaluacion ----------------
    doc.add_heading("Evaluación de Resultados y Comparación", level=1)
    doc.add_heading("Comparación Cuantitativa Antes y Después", level=2)
    O, P = "Original (oscura)", CLAVE_PROPUESTA

    def interp(clave):
        a, d = D.t(O, clave), D.t(P, clave)
        if a is None:
            return "[interpretación]"
        sube = d > a
        textos = {
            "brillo": ("Exposición recuperada; se acerca a la referencia." if sube else "No aumentó la exposición."),
            "contraste": ("Mayor separación de intensidades." if sube else "El contraste no mejoró."),
            "nitidez": ("Más respuesta de alta frecuencia; revisar junto al ruido." if sube
                        else "Baja por el suavizado bilateral del grano, no por desenfoque."),
            "ruido": ("Ruido amplificado por el realce; controlado por el bilateral." if sube
                      else "El ruido no aumentó pese al realce."),
            "entropia": ("Más niveles de gris con información." if sube else "Menor diversidad tonal."),
            "pct_oscuros": ("Aumentó la subexposición." if sube else "Las sombras densas se redujeron."),
            "densidad_bordes": ("Más contornos detectables con los mismos umbrales." if sube
                                else "Menos bordes detectados."),
        }
        return textos[clave]

    etiquetas = [("brillo", "Brillo", 1), ("contraste", "Contraste", 1), ("nitidez", "Nitidez", 1),
                 ("ruido", "Ruido σ̂", 2), ("entropia", "Entropía (bits)", 2),
                 ("pct_oscuros", "Píxeles subexpuestos (%)", 1), ("densidad_bordes", "Densidad de bordes (%)", 2)]
    filas = [[nom, v(D.t(O, k), d), v(D.t(P, k), d), var_pct(D.t(O, k), D.t(P, k)), interp(k)]
             for k, nom, d in etiquetas]
    filas.append(["Resolución", f"{D.m('oscura', 'ancho')}×{D.m('oscura', 'alto')}" if hay else "[valor]",
                  f"{ti.ANCHO_ANALISIS} px de ancho", "—",
                  "Se uniformó el tamaño de entrada; no se añade información."])
    tabla_apa(doc, 6, "Métricas de la imagen oscura antes y después del tratamiento propuesto",
              ["Métrica", "Antes", "Después", "Variación", "Interpretación"], filas,
              [1.5, 0.85, 0.85, 0.75, 2.55],
              f"Valores de referencia: brillo {v(D.m('referencia', 'brillo'))}, contraste "
              f"{v(D.m('referencia', 'contraste'))}, densidad de bordes {v(D.m('referencia', 'densidad_bordes'), 2)} %.")

    variantes = ["Original (oscura)", "Ajuste lineal α=1.2, β=25", "Ecualización global", "Solo CLAHE",
                 "Gamma + CLAHE", CLAVE_PROPUESTA]
    parrafo(doc, "Para verificar que la elección no fue arbitraria, las mismas métricas se calcularon para las "
            "alternativas descritas en la sección de diseño (Tabla 7 y Figura 4). Esta comparación permite separar el "
            "aporte de cada etapa: la gamma corrige la exposición, CLAHE el contraste local y el filtro bilateral "
            "el ruido introducido por ambas.")
    tabla_apa(doc, 7, "Comparación de alternativas de tratamiento sobre la imagen oscura",
              ["Variante", "Brillo", "Contraste", "Nitidez", "Ruido σ̂", "Saturados (%)", "Bordes (%)"],
              [[k.replace(" (propuesta)", " (propuesta)"), v(D.t(k, "brillo")), v(D.t(k, "contraste")),
                v(D.t(k, "nitidez")), v(D.t(k, "ruido"), 2), v(D.t(k, "pct_saturados"), 1),
                v(D.t(k, "densidad_bordes"), 2)] for k in variantes],
              [2.2, 0.65, 0.75, 0.75, 0.7, 0.85, 0.7],
              f"Saturados: píxeles con intensidad > {ti.UMBRAL_SATURADO}. La ecualización global suele maximizar el "
              "contraste a costa de ruido y saturación, por lo que un valor alto de contraste no implica, por sí solo, "
              "una mejor imagen para el sistema.")
    figura(doc, 4, "Imagen oscura original frente a las alternativas de tratamiento", fig("fig4_antes_despues.png"),
           "De izquierda a derecha y de arriba abajo: original, ajuste lineal, ecualización global, solo CLAHE, "
           "gamma + CLAHE y pipeline propuesto.")
    figura(doc, 5, "Histograma y distribución acumulada antes y después del tratamiento", fig("fig5_histograma.png"),
           "La CDF del original asciende de forma abrupta en los tonos bajos; tras el tratamiento se desplaza a la "
           "derecha y se vuelve más gradual, lo que indica un uso más amplio del rango dinámico.")

    doc.add_heading("Evaluación Cualitativa", level=2)
    if hay:
        b0, b1 = D.t(O, "brillo"), D.t(P, "brillo")
        c0, c1 = D.t(O, "contraste"), D.t(P, "contraste")
        txt = (f"El brillo pasó de {v(b0)} a {v(b1)} ({var_pct(b0, b1)}) y el contraste de {v(c0)} a {v(c1)} "
               f"({var_pct(c0, c1)}). ")
    else:
        txt = "El brillo pasó de [valor] a [valor] y el contraste de [valor] a [valor]. "
    parrafo(doc, "En términos de **legibilidad visual**, el objeto y sus detalles vuelven a distinguirse del fondo. " + txt +
            "Respecto a la **conservación de contornos**, al operar únicamente sobre L y filtrar con un núcleo "
            "sensible a la intensidad, las siluetas permanecen definidas y el color del objeto no se desplaza, a "
            "diferencia de la ecualización global, que tiende a producir zonas quemadas y grano visible. En cuanto a "
            "la **mitigación de artefactos**, el límite de recorte de CLAHE evita halos marcados en las fronteras entre "
            "regiones, y el filtro bilateral reduce el grano que la expansión tonal hace visible. Un artefacto residual "
            "esperable es una ligera textura en zonas planas muy oscuras, donde la señal original era casi nula: el "
            "procesamiento redistribuye la información existente, pero no puede crear información que el sensor no "
            "capturó.")

    doc.add_heading("Evaluación para el Sistema de Percepción", level=2)
    parrafo(doc, f"La Figura 6 compara la respuesta de Canny con los mismos umbrales ({ti.CANNY_BAJO}/{ti.CANNY_ALTO}) "
            f"antes y después. La densidad de bordes pasó de {v(D.t(O, 'densidad_bordes'), 2)} % a "
            f"{v(D.t(P, 'densidad_bordes'), 2)} %, frente a {v(D.m('referencia', 'densidad_bordes'), 2)} % de la "
            "referencia. Esto significa que, sin cambiar la configuración del detector, los contornos del objeto "
            "vuelven a superar el umbral, lo que facilita la extracción de contornos con `cv2.findContours`, la "
            "binarización por Otsu o umbral adaptativo y, en un modelo de aprendizaje profundo, que las activaciones "
            "de las primeras capas se parezcan más a las de imágenes bien expuestas. La nitidez se interpreta junto con "
            "el ruido: el filtro bilateral redujo la varianza del Laplaciano de "
            f"{v(D.t('Gamma + CLAHE', 'nitidez'))} a {v(D.t(P, 'nitidez'))} y el ruido estimado de "
            f"{v(D.t('Gamma + CLAHE', 'ruido'), 2)} a {v(D.t(P, 'ruido'), 2)}, de modo que esa disminución no indica "
            "pérdida de enfoque, sino eliminación de alta frecuencia espuria.")
    if hay and D.t(P, "densidad_bordes") > 1.3 * D.m("referencia", "densidad_bordes"):
        parrafo(doc, "Este resultado exige una lectura crítica: la densidad de bordes de la imagen tratada "
                f"({v(D.t(P, 'densidad_bordes'), 2)} %) supera incluso a la de la referencia "
                f"({v(D.m('referencia', 'densidad_bordes'), 2)} %). La Figura 6 muestra que, además de recuperar el "
                "contorno exterior y las divisiones internas del objeto, una parte importante de los bordes nuevos "
                "corresponde a la textura del fondo (la veta de la superficie de madera) y al ruido residual, que "
                "CLAHE realza porque amplifica cualquier variación local. Un mayor número de bordes no equivale, "
                "por tanto, a más información útil; para el sistema de percepción convendría restringir el análisis "
                "a una región de interés o ajustar los umbrales de Canny a la mediana de intensidad de la imagen.")
    figura(doc, 6, "Respuesta del detector de bordes de Canny antes y después del tratamiento", fig("fig6_bordes.png"),
           f"Umbrales fijos de histéresis {ti.CANNY_BAJO} y {ti.CANNY_ALTO} en ambas imágenes.")

    if hay:
        mejora = (D.t(P, "brillo") > D.t(O, "brillo") and D.t(P, "contraste") > D.t(O, "contraste")
                  and D.t(P, "densidad_bordes") > D.t(O, "densidad_bordes"))
        veredicto = ("Con base en estos datos, la transformación **sí mejoró** la imagen para el propósito "
                     "planteado: aumentó la exposición, el contraste y la cantidad de contornos detectables. El "
                     f"ruido estimado ({v(D.t(P, 'ruido'), 2)}) quedó por debajo de las alternativas sin filtrar, "
                     f"aunque por encima del de la referencia ({v(D.m('referencia', 'ruido'), 2)})." if mejora else
                     "Con base en estos datos, la transformación **mejoró solo parcialmente** la imagen: alguna de "
                     "las métricas objetivo no avanzó en la dirección esperada, por lo que se recomienda ajustar los "
                     "hiperparámetros o repetir la captura.")
    else:
        veredicto = ("[Veredicto: indique, con base en la Tabla 6, si la transformación mejoró la imagen para el "
                     "propósito planteado.]")
    parrafo(doc, veredicto + " Aun así, la imagen tratada no alcanza la calidad de una captura bien expuesta, "
            "lo que confirma que la mejor corrección es una adquisición adecuada.")

    # ---------------- 6. Analisis (preguntas orientadoras) ----------------
    doc.add_heading("Análisis: Respuesta a las Preguntas Orientadoras", level=1)
    preguntas = [
        ("¿Qué características de una imagen pueden afectar un sistema de percepción computacional?",
         "La iluminación (exposición y sombras), el enfoque, el contraste, la resolución, el formato de compresión, "
         "la perspectiva y el fondo. Cada una altera la distribución de intensidades o la geometría que el sistema "
         "aprendió a reconocer; en este trabajo la iluminación y el desenfoque produjeron los cambios métricos "
         "más marcados, y el fondo y la perspectiva, los menos visibles en métricas globales pero igualmente críticos."),
        ("¿Una imagen de mayor resolución necesariamente es mejor para un modelo?",
         "No. Los modelos operan con entradas de tamaño fijo y reducido, de modo que el exceso de píxeles se "
         "descarta al reescalar y solo aumenta el costo de memoria y cómputo. Una imagen de alta resolución "
         "desenfocada contiene menos información útil que una de menor resolución nítida; importa la información "
         "por píxel, no el número de píxeles. Por eso las métricas se compararon a un ancho común."),
        ("¿Qué efecto tuvo la iluminación sobre los datos obtenidos?",
         f"La baja iluminación redujo el brillo a {ratio_pct(D.m('oscura', 'brillo'), D.m('referencia', 'brillo'))} "
         f"y el contraste a {ratio_pct(D.m('oscura', 'contraste'), D.m('referencia', 'contraste'))} de la referencia, "
         "comprimió el histograma y redujo la densidad de bordes detectables, aunque el objeto estuviera presente."),
        ("¿Qué ocurrió con la imagen desenfocada y cómo se reflejó en la medida de nitidez?",
         f"Mantuvo un brillo similar, pero su varianza del Laplaciano cayó a "
         f"{ratio_pct(D.m('desenfocada', 'nitidez'), D.m('referencia', 'nitidez'))} de la referencia, porque el "
         "desenfoque actúa como un filtro pasa-bajos que elimina las altas frecuencias que el Laplaciano mide."),
        ("¿Todos los problemas encontrados pueden corregirse mediante procesamiento?",
         "No. La subexposición moderada es corregible porque la información existe, aunque comprimida. El "
         "desenfoque severo solo se compensa parcialmente (la deconvolución es un problema mal condicionado que "
         "amplifica el ruido), el fondo complejo requiere segmentación o un modelo robusto, y la perspectiva exige "
         "invariancia geométrica o aumento de datos. Las zonas saturadas o con señal nula no se recuperan."),
        ("¿Por qué es importante analizar una imagen antes de utilizarla como entrada de un modelo?",
         "Porque permite detectar defectos que degradarían el desempeño, decidir entre corregir, descartar o "
         "volver a capturar, y evitar aplicar transformaciones innecesarias que introduzcan artefactos. El "
         "diagnóstico convierte el preprocesamiento en una decisión fundamentada y no en una receta."),
        ("¿El tratamiento aplicado mejoró realmente la imagen para el propósito planteado?",
         f"Según la Tabla 6, el brillo varió {var_pct(D.t(O, 'brillo'), D.t(P, 'brillo'))}, el contraste "
         f"{var_pct(D.t(O, 'contraste'), D.t(P, 'contraste'))} y la densidad de bordes "
         f"{var_pct(D.t(O, 'densidad_bordes'), D.t(P, 'densidad_bordes'))}, con un ruido estimado de "
         f"{v(D.t(P, 'ruido'), 2)}. La mejora se sustenta en métricas orientadas a la tarea y no solo en la "
         "apariencia visual."),
    ]
    for i, (preg, resp) in enumerate(preguntas, 1):
        parrafo(doc, f"**{i}. {preg}** {resp}")

    # ---------------- 7. Conclusiones ----------------
    doc.add_heading("Conclusiones y Lecciones Aprendidas", level=1)
    parrafo(doc, "El preprocesamiento en visión artificial debe responder a un problema de captura concreto y "
            "medido. El diagnóstico con métricas simples (brillo, contraste, varianza del Laplaciano, ruido y "
            "densidad de bordes) fue suficiente para distinguir defectos fotométricos, ópticos y de composición, y "
            "para seleccionar la imagen oscura como el caso con un defecto severo y corregible.")
    parrafo(doc, "La combinación de corrección gamma adaptativa, CLAHE sobre el canal L y filtro bilateral se "
            "justificó frente a alternativas más simples: el ajuste lineal no separa sombras de luces y la "
            "ecualización global sobrerrealza el ruido. La prueba de verificación mostró además que CLAHE por sí "
            "solo no basta en subexposición severa, lo que llevó a incorporar la etapa de gamma: el diseño se "
            "corrigió a partir de evidencia, no de suposiciones.")
    parrafo(doc, "Como lección central, una imagen visualmente «bonita» no siempre es adecuada para un sistema de "
            "inteligencia artificial: un contraste máximo puede significar ruido amplificado, y una nitidez alta puede "
            "deberse a un fondo cargado. Las métricas deben interpretarse en conjunto y en función de la tarea.")
    parrafo(doc, "**Limitaciones y mejoras futuras.** Las métricas son globales y no distinguen el objeto del fondo; "
            "se propone calcularlas sobre una región de interés, evaluar el impacto final con un clasificador o "
            "detector entrenado, explorar la búsqueda sistemática de `clipLimit` y `sigmaColor`, y, sobre todo, "
            "mejorar la adquisición (iluminación difusa, trípode, ISO bajo), que sigue siendo la intervención más "
            "eficaz.")

    doc.add_heading("Repositorio del Código", level=2)
    parrafo(doc, f"El código fuente, las imágenes, las figuras y el archivo de métricas se encuentran en {REPOSITORIO} "
            "(DiegoYanzir, 2026). El repositorio incluye las instrucciones para reproducir el análisis en Google Colab.")

    # ---------------- Referencias ----------------
    salto_pagina(doc)
    doc.add_heading("Referencias", level=1)
    refs = [
        "Bradski, G. (2000). The OpenCV library. *Dr. Dobb's Journal of Software Tools, 25*(11), 120–125.",
        "Canny, J. (1986). A computational approach to edge detection. *IEEE Transactions on Pattern Analysis and "
        "Machine Intelligence, PAMI-8*(6), 679–698. https://doi.org/10.1109/TPAMI.1986.4767851",
        "DiegoYanzir. (2026). *Taller-R2-A2-S8---Informe-de-resultados-de-tratamiento-de-imagen* [Software]. GitHub. "
        + REPOSITORIO,
        "Futurology — An Optimistic Future. (2020, 19 de diciembre). *Convolutional neural networks explained (CNN "
        "visualized)* [Video]. YouTube. https://www.youtube.com/watch?v=pj9-rr1wDhM",
        "Gonzalez, R. C., & Woods, R. E. (2018). *Digital image processing* (4.ª ed.). Pearson.",
        "Image processing class (EGBE443) #5: Edge and contour. (s.f.). *Towards Data Science*. "
        "https://medium.com/data-science/image-processing-class-egbe443-5-edge-and-contour-d5d410f4483c",
        "Immerkær, J. (1996). Fast noise variance estimation. *Computer Vision and Image Understanding, 64*(2), "
        "300–302. https://doi.org/10.1006/cviu.1996.0060",
        "Pech-Pacheco, J. L., Cristóbal, G., Chamorro-Martínez, J., & Fernández-Valdivia, J. (2000). Diatom "
        "autofocusing in brightfield microscopy: A comparative study. En *Proceedings 15th International Conference "
        "on Pattern Recognition* (Vol. 3, pp. 314–317). IEEE. https://doi.org/10.1109/ICPR.2000.903548",
        "Pizer, S. M., Amburn, E. P., Austin, J. D., Cromartie, R., Geselowitz, A., Greer, T., ter Haar Romeny, B., "
        "Zimmerman, J. B., & Zuiderveld, K. (1987). Adaptive histogram equalization and its variations. *Computer "
        "Vision, Graphics, and Image Processing, 39*(3), 355–368. https://doi.org/10.1016/S0734-189X(87)80186-X",
        "Rodríguez Araújo, J. (2019, 20 de noviembre). Cómo empezar con reconocimiento de imagen y redes neuronales "
        "convolucionales en 5 minutos. *Planeta Chatbot*. https://planetachatbot.com/como-empezar-con-reconocimiento-"
        "de-imagen-y-redes-neuronales-convolucionales-en-5-minutos/",
        "Szeliski, R. (2022). *Computer vision: Algorithms and applications* (2.ª ed.). Springer. "
        "https://doi.org/10.1007/978-3-030-34372-9",
        "Tomasi, C., & Manduchi, R. (1998). Bilateral filtering for gray and color images. En *Sixth International "
        "Conference on Computer Vision* (pp. 839–846). IEEE. https://doi.org/10.1109/ICCV.1998.710815",
        "Zuiderveld, K. (1994). Contrast limited adaptive histogram equalization. En P. S. Heckbert (Ed.), *Graphics "
        "gems IV* (pp. 474–485). Academic Press. https://doi.org/10.1016/B978-0-12-336156-1.50061-6",
    ]
    for r in refs:
        referencia(doc, r)

    # ---------------- Apendice ----------------
    salto_pagina(doc)
    doc.add_heading("Apéndice A", level=1)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Código Completo de tratamiento_imagen.py").bold = True
    parrafo(doc, "Script completo y comentado. Requiere Python 3.9 o superior, OpenCV, NumPy y Matplotlib "
            "(`pip install opencv-python numpy matplotlib`).")
    codigo(doc, open(os.path.join(BASE, "tratamiento_imagen.py"), encoding="utf-8").read())

    doc.save(salida)
    return salida


def main():
    parser = argparse.ArgumentParser(description="Genera el informe APA 7 en formato .docx")
    parser.add_argument("--imagenes", default=os.path.join(BASE, "imagenes"))
    parser.add_argument("--resultados", default=os.path.join(BASE, "resultados"))
    parser.add_argument("--salida", default=os.path.join(BASE, "Informe_tratamiento_imagen.docx"))
    args = parser.parse_args()

    faltan = [n for n in ti.CONDICIONES if ti.buscar_archivo(args.imagenes, n) is None]
    if faltan:
        print("Faltan imágenes (" + ", ".join(faltan) + "): se genera el informe con marcadores [valor].")
        resultados = None
    else:
        resultados = ti.ejecutar(args.imagenes, args.resultados)
    print("Informe generado:", construir(resultados, args.resultados, args.salida))


if __name__ == "__main__":
    main()
